from pptx import Presentation
import os

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

def replace_image(shape, image_path):
    # This only works on picture shapes
    if shape.shape_type != 13:  # 13 = picture
        return False

    # Get the relationship ID of the image
    rId = shape._element.blipFill.blip.rEmbed

    # Replace the part with a new image
    image_part = shape.part.related_part(rId)
    image_part._blob = open(image_path, "rb").read()

    return True



prs = Presentation("reference.pptx")

slides ={}
for slide_index, slide in enumerate(prs.slides, start=1):
    print(f"Slide {slide_index} ")

    for shape in slide.shapes:
        slides[shape.name]=shape



for key,val in slides.items():
    print(key)
    for val in val.shapes:
        print(val.name)
        
        
